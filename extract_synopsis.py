from docx import Document

doc = Document('Synopsis_Main.docx')
print("=== PARAGRAPHS ===")
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text)

print("\n=== TABLES ===")
for ti, table in enumerate(doc.tables):
    print(f"\nTable {ti+1}:")
    for row in table.rows:
        print(" | ".join([cell.text for cell in row.cells]))
